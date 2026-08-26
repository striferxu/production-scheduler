"""生产排程系统 - FastAPI 主应用"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from fastapi import FastAPI, Depends, HTTPException, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from sqlalchemy.orm import Session
from datetime import datetime, timedelta
from typing import List, Optional
import json, io

from backend.models.database import engine, SessionLocal, get_db, Base
from backend.models.all_models import *
from backend.services.scheduler_engine import SchedulerEngine
from backend.auth import create_token, get_current_user, require_admin

# 创建所有表
Base.metadata.create_all(bind=engine)

app = FastAPI(title="生产排程系统", version="1.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ========== 静态文件 ==========
FRONTEND_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'frontend')
if os.path.exists(FRONTEND_DIR):
    app.mount("/static", StaticFiles(directory=FRONTEND_DIR), name="static")

EXPORTS_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'exports')
os.makedirs(EXPORTS_DIR, exist_ok=True)

# ========== 认证 ==========
from passlib.context import CryptContext
pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")

def verify_password(plain, hashed):
    return pwd_context.verify(plain, hashed)

def get_password_hash(password):
    return pwd_context.hash(password)

# ========== 首页 ==========
@app.get("/")
async def root():
    index_path = os.path.join(FRONTEND_DIR, 'index.html')
    if os.path.exists(index_path):
        return FileResponse(index_path)
    return {"message": "生产排程系统 API v1.0", "docs": "/docs"}

# ========== 用户认证 ==========
@app.post("/api/auth/login")
async def login(data: dict, db: Session = Depends(get_db)):
    user = db.query(User).filter(User.username == data.get('username')).first()
    if not user or not verify_password(data.get('password', ''), user.password_hash):
        raise HTTPException(status_code=401, detail="用户名或密码错误")
    return {"id": user.id, "username": user.username, "role": user.role,
            "token": create_token(user)}

@app.post("/api/auth/register")
async def register(data: dict, db: Session = Depends(get_db),
                   _auth: User = Depends(require_admin)):
    """仅管理员可创建用户账号；角色统一按入参，默认普通用户，不允许越级提权。"""
    role = data.get('role', 'user')
    if role not in ('admin', 'user'):
        raise HTTPException(status_code=400, detail="角色不合法")
    if db.query(User).filter(User.username == data.get('username')).first():
        raise HTTPException(status_code=400, detail="用户名已存在")
    user = User(username=data['username'],
                password_hash=get_password_hash(data.get('password') or ''), role=role)
    db.add(user)
    db.commit()
    return {"id": user.id, "username": user.username, "role": user.role}

# ========== CRUD 通用辅助 ==========
def crud_list(model, db, order_by=None):
    q = db.query(model)
    if order_by is not None:
        q = q.order_by(order_by)
    return q.all()

def crud_get(model, db, id):
    obj = db.query(model).get(id)
    if not obj:
        raise HTTPException(status_code=404, detail=f"{model.__tablename__} not found")
    return obj

def crud_create(model, db, data: dict):
    obj = model(**{k: v for k, v in data.items() if hasattr(model, k)})
    db.add(obj)
    db.commit()
    db.refresh(obj)
    return obj

def crud_update(model, db, id, data: dict):
    obj = crud_get(model, db, id)
    for k, v in data.items():
        if hasattr(obj, k) and k != 'id':
            setattr(obj, k, v)
    db.commit()
    return obj

def crud_delete(model, db, id):
    obj = crud_get(model, db, id)
    db.delete(obj)
    db.commit()
    return {"ok": True}

def model_to_dict(obj):
    if obj is None:
        return None
    result = {}
    for c in obj.__table__.columns:
        val = getattr(obj, c.name)
        if isinstance(val, datetime):
            val = val.isoformat()
        result[c.name] = val
    return result


def delete_guard(db, entity, id):
    """删除前引用完整性检查，被引用的数据禁止删除。"""
    def _ref(model, col, desc):
        if db.query(model).filter(col == id).first():
            raise HTTPException(status_code=400,
                                detail=f"无法删除：该{desc}仍被关联数据引用，请先解除关联（含历史排程任务）")

    if entity == 'product-categories':
        _ref(Product, Product.category_id, '产品分类')
    elif entity == 'products':
        _ref(ProductRoute, ProductRoute.product_id, '产品')
        _ref(ScheduleTask, ScheduleTask.product_id, '产品')
    elif entity == 'device-groups':
        _ref(Device, Device.group_id, '设备组')
    elif entity == 'devices':
        _ref(process_device, process_device.c.device_id, '设备')
        _ref(employee_device, employee_device.c.device_id, '设备')
        _ref(ScheduleTask, ScheduleTask.device_id, '设备')
    elif entity == 'processes':
        _ref(RouteStep, RouteStep.process_id, '工序')
        _ref(employee_process, employee_process.c.process_id, '工序')
        _ref(process_device, process_device.c.process_id, '工序')
        _ref(process_device_group, process_device_group.c.process_id, '工序')
        _ref(ScheduleTask, ScheduleTask.process_id, '工序')
    elif entity == 'product-routes':
        step_ids = [rs.id for rs in db.query(RouteStep).filter(RouteStep.route_id == id).all()]
        if step_ids:
            if db.query(ScheduleTask).filter(ScheduleTask.route_step_id.in_(step_ids)).first():
                raise HTTPException(status_code=400,
                                    detail="无法删除：该工艺路线仍有排程任务引用其工序步骤，请先解除（含历史排程任务）")
    elif entity == 'shifts':
        _ref(Employee, Employee.shift_id, '班次')
        _ref(shift_rotation, shift_rotation.c.shift_id, '班次')
    elif entity == 'rotations':
        _ref(Employee, Employee.rotation_id, '轮班模式')
        _ref(shift_rotation, shift_rotation.c.rotation_id, '轮班模式')
    elif entity == 'employee-groups':
        _ref(Employee, Employee.group_id, '人员组别')
    elif entity == 'employees':
        _ref(ScheduleTask, ScheduleTask.employee_id, '员工')
    elif entity == 'users':
        # 禁止删除当前登录的最后一个管理员，避免系统失去管理
        target = db.query(User).get(id)
        if target and target.role == 'admin':
            admin_count = db.query(User).filter(User.role == 'admin').count()
            if admin_count <= 1:
                raise HTTPException(status_code=400, detail="无法删除：系统至少需保留一个管理员账号")
    return True


def validate_schedule_input(data: dict):
    """校验排程入参：产品、数量、交期合法性。"""
    products = data.get('products', [])
    if not products:
        raise HTTPException(status_code=400, detail="请至少选择一个产品")
    seen = set()
    for i, pd in enumerate(products, 1):
        pid = pd.get('product_id')
        qty = pd.get('quantity')
        if not pid:
            raise HTTPException(status_code=400, detail=f"第{i}个产品缺少产品ID")
        try:
            qty = float(qty)
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail=f"第{i}个产品数量无效")
        if qty <= 0:
            raise HTTPException(status_code=400, detail=f"第{i}个产品数量必须大于0")
        if pid in seen:
            raise HTTPException(status_code=400, detail=f"产品 {pid} 重复，请合并数量")
        seen.add(pid)
        due = pd.get('due_date')
        if due:
            try:
                datetime.fromisoformat(str(due).replace('Z', '+00:00'))
            except ValueError:
                raise HTTPException(status_code=400, detail=f"第{i}个产品交期格式无效")


def schedule_overdue_report(tasks) -> list:
    """返回交期不达标（超期）的产品列表，供前端警示。"""
    if not tasks:
        return []
    overdue = {}
    for t in tasks:
        if t.due_date and t.end_time and t.end_time > t.due_date:
            o = overdue.setdefault(t.product_name,
                                   {"product": t.product_name, "due": t.due_date.isoformat(),
                                    "latest_end": None})
            if o["latest_end"] is None or t.end_time > datetime.fromisoformat(o["latest_end"]):
                o["latest_end"] = t.end_time.isoformat()
    return list(overdue.values())


def validate_route_steps(steps_data, title="工艺路线"):
    """校验工艺路线工序步骤：工时必须大于0。"""
    if not steps_data:
        raise HTTPException(status_code=400, detail=f"{title}至少需包含一道工序")
    for i, step in enumerate(steps_data, 1):
        if not step.get('process_id'):
            raise HTTPException(status_code=400, detail=f"{title}第{i}道工序未选择工序")
        try:
            t = float(step.get('process_time_per_unit', 0))
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail=f"{title}第{i}道工序单件工时无效")
        if t <= 0:
            raise HTTPException(status_code=400, detail=f"{title}第{i}道工序单件工时必须大于0")

# ========== 产品分类 ==========
@app.get("/api/product-categories")
def list_product_categories(db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    return [model_to_dict(x) for x in crud_list(ProductCategory, db, ProductCategory.name)]

@app.post("/api/product-categories")
def create_product_category(data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    return model_to_dict(crud_create(ProductCategory, db, data))

@app.put("/api/product-categories/{id}")
def update_product_category(id: int, data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    return model_to_dict(crud_update(ProductCategory, db, id, data))

@app.delete("/api/product-categories/{id}")
def delete_product_category(id: int, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    delete_guard(db, 'product-categories', id)
    return crud_delete(ProductCategory, db, id)

# ========== 产品 ==========
@app.get("/api/products")
def list_products(db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    products = db.query(Product).order_by(Product.name).all()
    return [model_to_dict(p) for p in products]

@app.post("/api/products")
def create_product(data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    return model_to_dict(crud_create(Product, db, data))

@app.put("/api/products/{id}")
def update_product(id: int, data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    return model_to_dict(crud_update(Product, db, id, data))

@app.delete("/api/products/{id}")
def delete_product(id: int, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    delete_guard(db, 'products', id)
    return crud_delete(Product, db, id)

# ========== 设备组 ==========
@app.get("/api/device-groups")
def list_device_groups(db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    return [model_to_dict(x) for x in crud_list(DeviceGroup, db, DeviceGroup.name)]

@app.post("/api/device-groups")
def create_device_group(data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    return model_to_dict(crud_create(DeviceGroup, db, data))

@app.put("/api/device-groups/{id}")
def update_device_group(id: int, data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    return model_to_dict(crud_update(DeviceGroup, db, id, data))

@app.delete("/api/device-groups/{id}")
def delete_device_group(id: int, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    delete_guard(db, 'device-groups', id)
    return crud_delete(DeviceGroup, db, id)

# ========== 设备 ==========
@app.get("/api/devices")
def list_devices(db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    devices = db.query(Device).order_by(Device.name).all()
    return [{**model_to_dict(d), 'group_name': d.group.name if d.group else None} for d in devices]

@app.post("/api/devices")
def create_device(data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    return model_to_dict(crud_create(Device, db, data))

@app.put("/api/devices/{id}")
def update_device(id: int, data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    return model_to_dict(crud_update(Device, db, id, data))

@app.delete("/api/devices/{id}")
def delete_device(id: int, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    delete_guard(db, 'devices', id)
    return crud_delete(Device, db, id)

# ========== 工序 ==========
@app.get("/api/processes")
def list_processes(db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    processes = db.query(Process).order_by(Process.name).all()
    result = []
    for p in processes:
        d = model_to_dict(p)
        d['device_ids'] = [dev.id for dev in p.devices]
        d['device_group_ids'] = [dg.id for dg in p.device_groups]
        result.append(d)
    return result

@app.post("/api/processes")
def create_process(data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    device_ids = data.pop('device_ids', [])
    device_group_ids = data.pop('device_group_ids', [])
    obj = Process(**{k: v for k, v in data.items() if hasattr(Process, k)})
    db.add(obj)
    db.flush()
    for did in device_ids:
        db.execute(process_device.insert().values(process_id=obj.id, device_id=did))
    for dgid in device_group_ids:
        db.execute(process_device_group.insert().values(process_id=obj.id, device_group_id=dgid))
    db.commit()
    return model_to_dict(obj)

@app.put("/api/processes/{id}")
def update_process(id: int, data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    obj = crud_get(Process, db, id)
    device_ids = data.pop('device_ids', None)
    device_group_ids = data.pop('device_group_ids', None)
    for k, v in data.items():
        if hasattr(obj, k) and k != 'id':
            setattr(obj, k, v)
    if device_ids is not None:
        db.execute(process_device.delete().where(process_device.c.process_id == id))
        for did in device_ids:
            db.execute(process_device.insert().values(process_id=id, device_id=did))
    if device_group_ids is not None:
        db.execute(process_device_group.delete().where(process_device_group.c.process_id == id))
        for dgid in device_group_ids:
            db.execute(process_device_group.insert().values(process_id=id, device_group_id=dgid))
    db.commit()
    return model_to_dict(obj)

@app.delete("/api/processes/{id}")
def delete_process(id: int, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    delete_guard(db, 'processes', id)
    return crud_delete(Process, db, id)

# ========== 工艺路线 ==========
@app.get("/api/product-routes")
def list_routes(db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    routes = db.query(ProductRoute).all()
    result = []
    for r in routes:
        rd = model_to_dict(r)
        rd['product_name'] = r.product.name if r.product else None
        rd['steps'] = []
        for s in sorted(r.steps, key=lambda x: x.step_order):
            sd = model_to_dict(s)
            sd['process_name'] = s.process.name if s.process else None
            rd['steps'].append(sd)
        result.append(rd)
    return result

@app.post("/api/product-routes")
def create_route(data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    steps_data = data.pop('steps', [])
    validate_route_steps(steps_data)
    route = ProductRoute(**{k: v for k, v in data.items() if hasattr(ProductRoute, k)})
    db.add(route)
    db.flush()
    for i, step in enumerate(steps_data):
        rs = RouteStep(
            route_id=route.id,
            process_id=step['process_id'],
            step_order=step.get('step_order', i),
            process_time_per_unit=float(step.get('process_time_per_unit', 0))
        )
        db.add(rs)
    db.commit()
    return model_to_dict(route)

@app.put("/api/product-routes/{id}")
def update_route(id: int, data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    route = crud_get(ProductRoute, db, id)
    steps_data = data.pop('steps', None)
    for k, v in data.items():
        if hasattr(route, k) and k != 'id':
            setattr(route, k, v)
    if steps_data is not None:
        validate_route_steps(steps_data)
        db.query(RouteStep).filter(RouteStep.route_id == id).delete()
        for i, step in enumerate(steps_data):
            rs = RouteStep(
                route_id=id,
                process_id=step['process_id'],
                step_order=step.get('step_order', i),
                process_time_per_unit=float(step.get('process_time_per_unit', 0))
            )
            db.add(rs)
    db.commit()
    return model_to_dict(route)

@app.delete("/api/product-routes/{id}")
def delete_route(id: int, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    delete_guard(db, 'product-routes', id)
    db.query(RouteStep).filter(RouteStep.route_id == id).delete()
    return crud_delete(ProductRoute, db, id)

# ========== 班次 ==========
@app.get("/api/shifts")
def list_shifts(db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    return [model_to_dict(x) for x in crud_list(Shift, db, Shift.name)]

@app.post("/api/shifts")
def create_shift(data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    return model_to_dict(crud_create(Shift, db, data))

@app.put("/api/shifts/{id}")
def update_shift(id: int, data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    return model_to_dict(crud_update(Shift, db, id, data))

@app.delete("/api/shifts/{id}")
def delete_shift(id: int, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    delete_guard(db, 'shifts', id)
    return crud_delete(Shift, db, id)

# ========== 轮班模式 ==========
@app.get("/api/rotations")
def list_rotations(db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    rotations = db.query(Rotation).all()
    result = []
    for r in rotations:
        rd = model_to_dict(r)
        rd['shifts'] = [model_to_dict(s) for s in r.shift_rotations]
        result.append(rd)
    return result

@app.post("/api/rotations")
def create_rotation(data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    shift_ids = data.pop('shift_ids', [])
    obj = Rotation(**{k: v for k, v in data.items() if hasattr(Rotation, k)})
    db.add(obj)
    db.flush()
    for i, sid in enumerate(shift_ids):
        db.execute(shift_rotation.insert().values(
            rotation_id=obj.id, shift_id=sid, sequence_order=i))
    db.commit()
    return model_to_dict(obj)

@app.put("/api/rotations/{id}")
def update_rotation(id: int, data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    obj = crud_get(Rotation, db, id)
    shift_ids = data.pop('shift_ids', None)
    for k, v in data.items():
        if hasattr(obj, k) and k != 'id':
            setattr(obj, k, v)
    if shift_ids is not None:
        db.execute(shift_rotation.delete().where(shift_rotation.c.rotation_id == id))
        for i, sid in enumerate(shift_ids):
            db.execute(shift_rotation.insert().values(
                rotation_id=id, shift_id=sid, sequence_order=i))
    db.commit()
    return model_to_dict(obj)

@app.delete("/api/rotations/{id}")
def delete_rotation(id: int, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    delete_guard(db, 'rotations', id)
    return crud_delete(Rotation, db, id)

# ========== 人员组别 ==========
@app.get("/api/employee-groups")
def list_employee_groups(db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    return [model_to_dict(x) for x in crud_list(EmployeeGroup, db, EmployeeGroup.name)]

@app.post("/api/employee-groups")
def create_employee_group(data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    return model_to_dict(crud_create(EmployeeGroup, db, data))

@app.put("/api/employee-groups/{id}")
def update_employee_group(id: int, data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    return model_to_dict(crud_update(EmployeeGroup, db, id, data))

@app.delete("/api/employee-groups/{id}")
def delete_employee_group(id: int, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    delete_guard(db, 'employee-groups', id)
    return crud_delete(EmployeeGroup, db, id)

# ========== 员工 ==========
@app.get("/api/employees")
def list_employees(db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    employees = db.query(Employee).order_by(Employee.name).all()
    result = []
    for e in employees:
        ed = model_to_dict(e)
        ed['group_name'] = e.group.name if e.group else None
        ed['shift_name'] = e.shift.name if e.shift else None
        ed['rotation_name'] = e.rotation.name if e.rotation else None
        ed['device_ids'] = [d.id for d in e.devices]
        ed['process_ids'] = [p.id for p in e.processes]
        result.append(ed)
    return result

@app.post("/api/employees")
def create_employee(data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    device_ids = data.pop('device_ids', [])
    process_ids = data.pop('process_ids', [])
    obj = Employee(**{k: v for k, v in data.items() if hasattr(Employee, k)})
    db.add(obj)
    db.flush()
    for did in device_ids:
        db.execute(employee_device.insert().values(employee_id=obj.id, device_id=did))
    for pid in process_ids:
        db.execute(employee_process.insert().values(employee_id=obj.id, process_id=pid))
    db.commit()
    return model_to_dict(obj)

@app.put("/api/employees/{id}")
def update_employee(id: int, data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    obj = crud_get(Employee, db, id)
    device_ids = data.pop('device_ids', None)
    process_ids = data.pop('process_ids', None)
    for k, v in data.items():
        if hasattr(obj, k) and k != 'id':
            setattr(obj, k, v)
    if device_ids is not None:
        db.execute(employee_device.delete().where(employee_device.c.employee_id == id))
        for did in device_ids:
            db.execute(employee_device.insert().values(employee_id=id, device_id=did))
    if process_ids is not None:
        db.execute(employee_process.delete().where(employee_process.c.employee_id == id))
        for pid in process_ids:
            db.execute(employee_process.insert().values(employee_id=id, process_id=pid))
    db.commit()
    return model_to_dict(obj)

@app.delete("/api/employees/{id}")
def delete_employee(id: int, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    delete_guard(db, 'employees', id)
    return crud_delete(Employee, db, id)

# ========== 用户管理 ==========
@app.get("/api/users")
def list_users(db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    users = db.query(User).order_by(User.username).all()
    return [{k: v for k, v in model_to_dict(u).items() if k != 'password_hash'} for u in users]

@app.post("/api/users")
def create_user(data: dict, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    import os, secrets
    pwd = data.pop('password', None) or os.environ.get('DEFAULT_USER_PASSWORD') or secrets.token_urlsafe(10)
    data['password_hash'] = get_password_hash(pwd)
    return model_to_dict(crud_create(User, db, data))

@app.delete("/api/users/{id}")
def delete_user(id: int, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    delete_guard(db, 'users', id)
    return crud_delete(User, db, id)

# ========== 排程引擎 ==========
scheduler = SchedulerEngine()

@app.post("/api/schedule/run")
async def run_schedule(data: dict, db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    """执行排程
    输入: {products: [{product_id, quantity, due_date}], locked_task_ids: []}
    """
    try:
        validate_schedule_input(data)
        result = scheduler.run(db, data)
        overdue = schedule_overdue_report(result)
        return {"ok": True, "tasks": len(result),
                "batch_id": result[0].batch_id if result else None,
                "overdue": overdue}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.get("/api/schedule/tasks")
def get_schedule_tasks(batch_id: Optional[str] = None, db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    q = db.query(ScheduleTask)
    if batch_id:
        q = q.filter(ScheduleTask.batch_id == batch_id)
    else:
        # 获取最新批次
        latest = db.query(ScheduleTask.batch_id).order_by(ScheduleTask.created_at.desc()).first()
        if latest:
            q = q.filter(ScheduleTask.batch_id == latest[0])
    tasks = q.order_by(ScheduleTask.start_time).all()
    return [model_to_dict(t) for t in tasks]

@app.get("/api/schedule/batches")
def list_batches(db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    batches = db.query(ScheduleTask.batch_id, ScheduleTask.created_at)\
        .group_by(ScheduleTask.batch_id).order_by(ScheduleTask.created_at.desc()).all()
    return [{"batch_id": b[0], "created_at": b[1].isoformat() if b[1] else None} for b in batches]

@app.post("/api/schedule/tasks/{task_id}/lock")
def lock_task(task_id: int, data: dict, db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    """锁定/解锁任务"""
    task = crud_get(ScheduleTask, db, task_id)
    task.is_locked = data.get('is_locked', True)
    db.commit()
    return model_to_dict(task)

@app.post("/api/schedule/tasks/{task_id}/move")
def move_task(task_id: int, data: dict, db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    """拖拽移动任务：校验技能后更新分配并锁定，触发同批次联动重算。
    入参可选: {start_time, device_id, employee_id}"""
    task = crud_get(ScheduleTask, db, task_id)

    # 目标设备校验
    if data.get('device_id') is not None and data['device_id']:
        dev = db.query(Device).get(data['device_id'])
        if not dev:
            raise HTTPException(status_code=400, detail="目标设备不存在")
        if dev.status != 'available':
            raise HTTPException(status_code=400, detail="目标设备当前不可用（停用/保养中）")
        task.device_id = dev.id
        task.device_name = dev.name
        task.device_group_name = dev.group.name if dev.group else ''

    # 目标员工技能校验
    if data.get('employee_id') is not None and data['employee_id']:
        emp = db.query(Employee).get(data['employee_id'])
        if not emp:
            raise HTTPException(status_code=400, detail="目标员工不存在")
        has_proc = db.query(employee_process).filter(
            employee_process.c.employee_id == emp.id,
            employee_process.c.process_id == task.process_id).first()
        has_dev = True
        if task.device_id:
            has_dev = db.query(employee_device).filter(
                employee_device.c.employee_id == emp.id,
                employee_device.c.device_id == task.device_id).first()
        if not has_proc or not has_dev:
            raise HTTPException(status_code=400,
                                detail=f"技能不匹配：员工 {emp.name} 不具备工序「{task.process_name}」或该设备操作权限，无法放置")
        task.employee_id = emp.id
        task.employee_name = emp.name
        task.employee_group_name = emp.group.name if emp.group else ''
        task.shift_name = emp.shift.name if emp.shift else ''

    # 开始时间更新（支持 ISO 字符串或毫秒时间戳）
    if data.get('start_time') is not None:
        st = data['start_time']
        try:
            if isinstance(st, (int, float)):
                task.start_time = datetime.fromtimestamp(st / 1000.0)
            else:
                task.start_time = datetime.fromisoformat(str(st))
        except (ValueError, OSError, OverflowError):
            raise HTTPException(status_code=400, detail="开始时间格式无效")

    task.is_locked = True
    db.commit()

    # 联动重算：保留本任务(锁定)，对其余任务重新排程
    try:
        eng = SchedulerEngine()
        prods = {}
        for t in db.query(ScheduleTask).filter(ScheduleTask.batch_id == task.batch_id).all():
            if t.product_id:
                prods.setdefault(t.product_id, {
                    'product_id': t.product_id,
                    'quantity': t.quantity,
                    'due_date': (t.due_date or datetime.now()).isoformat(),
                })
        db.commit()
        eng.rerun(db, task.batch_id, list(prods.values()))
    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=f"移动后重算失败：{e}")

    return model_to_dict(task)

@app.post("/api/schedule/unlock-all")
def unlock_all(db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    db.query(ScheduleTask).update({ScheduleTask.is_locked: False})
    db.commit()
    return {"ok": True}

@app.post("/api/schedule/rerun")
async def rerun_schedule(data: dict, db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    """手工调整后重排"""
    batch_id = data.get('batch_id')
    products_data = data.get('products', [])
    try:
        validate_schedule_input({"products": products_data})
        result = scheduler.rerun(db, batch_id, products_data)
        return {"ok": True, "tasks": len(result)}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

# ========== Excel导出 ==========
@app.post("/api/export/excel")
def export_excel(data: dict, db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    """导出派工单 {batch_id, time_range: 'day'/'week'/'month'/custom, start_date, end_date}"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    
    batch_id = data.get('batch_id')
    time_range = data.get('time_range', 'day')
    start_date = data.get('start_date')
    end_date = data.get('end_date')
    
    q = db.query(ScheduleTask)
    if batch_id:
        q = q.filter(ScheduleTask.batch_id == batch_id)
    
    if time_range == 'day':
        today = datetime.now().strftime('%Y-%m-%d')
        q = q.filter(ScheduleTask.start_time >= f"{today} 00:00:00")
        q = q.filter(ScheduleTask.start_time <= f"{today} 23:59:59")
    elif time_range == 'week':
        today = datetime.now()
        monday = today - timedelta(days=today.weekday())
        sunday = monday + timedelta(days=6)
        q = q.filter(ScheduleTask.start_time >= monday.strftime('%Y-%m-%d 00:00:00'))
        q = q.filter(ScheduleTask.start_time <= sunday.strftime('%Y-%m-%d 23:59:59'))
    elif time_range == 'month':
        today = datetime.now()
        first_day = today.replace(day=1)
        if today.month == 12:
            last_day = today.replace(year=today.year+1, month=1, day=1) - timedelta(days=1)
        else:
            last_day = today.replace(month=today.month+1, day=1) - timedelta(days=1)
        q = q.filter(ScheduleTask.start_time >= first_day.strftime('%Y-%m-%d 00:00:00'))
        q = q.filter(ScheduleTask.start_time <= last_day.strftime('%Y-%m-%d 23:59:59'))
    elif time_range == 'custom' and start_date and end_date:
        q = q.filter(ScheduleTask.start_time >= f"{start_date} 00:00:00")
        q = q.filter(ScheduleTask.start_time <= f"{end_date} 23:59:59")
    
    tasks = q.order_by(ScheduleTask.start_time).all()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "派工单"
    
    # 表头样式
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font_white = Font(bold=True, size=11, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    headers = ['日期', '人员组别', '员工姓名', '班次', '设备组', '设备名称',
               '产品分类', '产品名称', '工序名称', '生产数量', '计划开始时间',
               '计划结束时间', '工时(分钟)', '交期']
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    
    for row, task in enumerate(tasks, 2):
        values = [
            task.start_time.strftime('%Y-%m-%d') if task.start_time else '',
            task.employee_group_name or '',
            task.employee_name or '',
            task.shift_name or '',
            task.device_group_name or '',
            task.device_name or '',
            task.product_category or '',
            task.product_name or '',
            task.process_name or '',
            task.quantity,
            task.start_time.strftime('%Y-%m-%d %H:%M') if task.start_time else '',
            task.end_time.strftime('%Y-%m-%d %H:%M') if task.end_time else '',
            round(task.duration_minutes, 1) if task.duration_minutes else '',
            task.due_date.strftime('%Y-%m-%d') if task.due_date else '',
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')
    
    # 调整列宽
    for col in range(1, len(headers)+1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 16
    
    filename = f"派工单_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(EXPORTS_DIR, filename)
    wb.save(filepath)
    return FileResponse(filepath, filename=filename, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# ========== 备份管理 ==========
@app.get("/api/backup")
def list_backups(db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    return [model_to_dict(b) for b in db.query(BackupRecord).order_by(BackupRecord.created_at.desc()).all()]

@app.post("/api/backup/create")
def create_backup(db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    import shutil, glob
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    backup_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'backups')
    os.makedirs(backup_dir, exist_ok=True)
    backup_file = os.path.join(backup_dir, f'backup_{timestamp}.db')
    db_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'data', 'scheduler.db')
    shutil.copy2(db_path, backup_file)
    record = BackupRecord(file_path=backup_file, backup_type='manual', file_size=os.path.getsize(backup_file))
    db.add(record)
    db.commit()
    return model_to_dict(record)

@app.post("/api/backup/restore/{record_id}")
def restore_backup(record_id: int, db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    import shutil
    record = crud_get(BackupRecord, db, record_id)
    db_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'data', 'scheduler.db')
    shutil.copy2(record.file_path, db_path)
    return {"ok": True, "message": "备份已恢复，请重启服务"}

# ========== Excel导入基础数据 ==========
@app.post("/api/import/excel")
async def import_excel(file: UploadFile = File(...), entity_type: str = "", db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    """从Excel导入基础数据"""
    from openpyxl import load_workbook
    content = await file.read()
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    errors = []
    imported = 0
    
    model_map = {
        'product_categories': ProductCategory,
        'products': Product,
        'device_groups': DeviceGroup,
        'devices': Device,
        'processes': Process,
        'employee_groups': EmployeeGroup,
        'employees': Employee,
        'shifts': Shift,
    }
    
    model = model_map.get(entity_type)
    if not model:
        raise HTTPException(400, f"未知实体类型: {entity_type}")
    
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        data = {}
        for i, header in enumerate(headers):
            if header and i < len(row):
                data[header] = row[i]
        try:
            obj = model(**{k: v for k, v in data.items() if hasattr(model, k)})
            db.add(obj)
            imported += 1
        except Exception as e:
            errors.append(f"第{row_idx}行: {str(e)}")
    
    db.commit()
    return {"imported": imported, "errors": errors}

@app.get("/api/export/excel-template/{entity_type}")
def export_template(entity_type: str, _auth: User = Depends(get_current_user)):
    """下载导入模板"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    
    templates = {
        'product_categories': ['name', 'description'],
        'products': ['name', 'category_id', 'description'],
        'device_groups': ['name', 'description', 'work_start_time', 'work_end_time', 'status'],
        'devices': ['name', 'group_id', 'description', 'work_start_time', 'work_end_time', 'status'],
        'processes': ['name', 'description', 'is_splittable'],
        'employee_groups': ['name', 'description'],
        'employees': ['name', 'group_id', 'employee_code'],
        'shifts': ['name', 'start_time', 'end_time', 'description'],
    }
    
    headers = templates.get(entity_type, ['name'])
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    
    filename = f"template_{entity_type}.xlsx"
    filepath = os.path.join(EXPORTS_DIR, filename)
    wb.save(filepath)
    return FileResponse(filepath, filename=filename, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# ========== 数据导出 ==========
@app.get("/api/export/all-data")
def export_all_data(db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    from openpyxl import Workbook
    wb = Workbook()
    
    entities = [
        ('产品分类', ProductCategory, ['id', 'name', 'description']),
        ('产品', Product, ['id', 'name', 'category_id', 'description']),
        ('设备组', DeviceGroup, ['id', 'name', 'description', 'work_start_time', 'work_end_time', 'status']),
        ('设备', Device, ['id', 'name', 'group_id', 'description', 'work_start_time', 'work_end_time', 'status']),
        ('工序', Process, ['id', 'name', 'description', 'is_splittable']),
        ('人员组别', EmployeeGroup, ['id', 'name', 'description']),
        ('员工', Employee, ['id', 'name', 'group_id', 'employee_code']),
        ('班次', Shift, ['id', 'name', 'start_time', 'end_time', 'description']),
    ]
    
    for sheet_name, model, columns in entities:
        ws = wb.create_sheet(sheet_name)
        for col, header in enumerate(columns, 1):
            ws.cell(row=1, column=col, value=header)
        for row, obj in enumerate(db.query(model).all(), 2):
            for col, col_name in enumerate(columns, 1):
                ws.cell(row=row, column=col, value=getattr(obj, col_name, ''))
    
    # 删除默认sheet
    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']
    
    filename = f"基础数据导出_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(EXPORTS_DIR, filename)
    wb.save(filepath)
    return FileResponse(filepath, filename=filename, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# ========== 获取各种选项 ==========
@app.get("/api/options/{entity_type}")
def get_options(entity_type: str, db: Session = Depends(get_db), _auth: User = Depends(get_current_user)):
    """获取下拉选项"""
    option_map = {
        'device_groups': (DeviceGroup, 'id', 'name'),
        'devices': (Device, 'id', 'name'),
        'processes': (Process, 'id', 'name'),
        'employee_groups': (EmployeeGroup, 'id', 'name'),
        'shifts': (Shift, 'id', 'name'),
        'rotations': (Rotation, 'id', 'name'),
        'product_categories': (ProductCategory, 'id', 'name'),
        'products': (Product, 'id', 'name'),
    }
    if entity_type not in option_map:
        raise HTTPException(404, f"未知选项类型: {entity_type}")
    model, id_field, name_field = option_map[entity_type]
    items = db.query(model).all()
    return [{"value": getattr(item, id_field), "label": getattr(item, name_field)} for item in items]




# ========== 示例数据（演示环境一键载入） ==========
@app.post("/api/seed/demo")
def seed_demo(db: Session = Depends(get_db), _auth: User = Depends(require_admin)):
    """仅当基础数据为空时，载入一套演示数据，便于快速体验排程。"""
    if db.query(Product).first() or db.query(Device).first() or db.query(Employee).first():
        raise HTTPException(status_code=400, detail="基础数据已存在，为避免覆盖，载入示例数据仅在空库时可用")
    from backend.services.scheduler_engine import SchedulerEngine  # noqa

    # 产品分类
    cat_machine = ProductCategory(name="冲压件", description="金属冲压类部件")
    cat_assemble = ProductCategory(name="装配件", description="需装配的组件")
    db.add_all([cat_machine, cat_assemble]); db.flush()

    # 产品
    p1 = Product(name="刹车片", category_id=cat_machine.id, description="前轮刹车片")
    p2 = Product(name="方向盘", category_id=cat_assemble.id, description="标准方向盘")
    p3 = Product(name="排气管", category_id=cat_machine.id, description="消音器排气管")
    db.add_all([p1, p2, p3]); db.flush()

    # 设备组（仅白天可用）
    g_stamp = DeviceGroup(name="冲压组", description="冲压设备", work_start_time="08:00", work_end_time="17:00", status="available")
    g_asm = DeviceGroup(name="装配组", description="装配设备", work_start_time="08:00", work_end_time="17:00", status="available")
    db.add_all([g_stamp, g_asm]); db.flush()

    # 设备
    d1 = Device(name="冲压机1号", group_id=g_stamp.id, status="available")
    d2 = Device(name="冲压机2号", group_id=g_stamp.id, status="available")
    d3 = Device(name="装配站1号", group_id=g_asm.id, status="available")
    d4 = Device(name="装配站2号", group_id=g_asm.id, status="available")
    db.add_all([d1, d2, d3, d4]); db.flush()

    # 工序（绑定设备 / 设备组）
    pro_stamp = Process(name="冲压成型", description="板材冲压", is_splittable=True)
    pro_drill = Process(name="钻孔", description="打孔", is_splittable=True)
    pro_asm = Process(name="装配", description="组装", is_splittable=True)
    db.add_all([pro_stamp, pro_drill, pro_asm]); db.flush()
    pro_stamp.device_groups = [g_stamp]
    pro_drill.devices = [d1, d2]
    pro_asm.device_groups = [g_asm]

    # 工艺路线
    r1 = ProductRoute(product_id=p1.id, name="刹车片标准工艺")
    r2 = ProductRoute(product_id=p2.id, name="方向盘标准工艺")
    r3 = ProductRoute(product_id=p3.id, name="排气管标准工艺")
    db.add_all([r1, r2, r3]); db.flush()
    db.add_all([
        RouteStep(route_id=r1.id, process_id=pro_stamp.id, step_order=0, process_time_per_unit=5),
        RouteStep(route_id=r1.id, process_id=pro_drill.id, step_order=1, process_time_per_unit=3),
        RouteStep(route_id=r2.id, process_id=pro_stamp.id, step_order=0, process_time_per_unit=3),
        RouteStep(route_id=r2.id, process_id=pro_asm.id, step_order=1, process_time_per_unit=8),
        RouteStep(route_id=r3.id, process_id=pro_stamp.id, step_order=0, process_time_per_unit=6),
        RouteStep(route_id=r3.id, process_id=pro_asm.id, step_order=1, process_time_per_unit=4),
    ])
    db.flush()

    # 班次
    sh_morning = Shift(name="早班", start_time="08:00", end_time="16:00")
    sh_mid = Shift(name="中班", start_time="16:00", end_time="00:00")
    sh_night = Shift(name="晚班", start_time="00:00", end_time="08:00")
    db.add_all([sh_morning, sh_mid, sh_night]); db.flush()

    # 轮班模式
    rot = Rotation(name="三班倒", cycle_days=3, description="早→中→晚循环")
    db.add(rot); db.flush()
    for i, sh in enumerate([sh_morning, sh_mid, sh_night]):
        db.execute(shift_rotation.insert().values(rotation_id=rot.id, shift_id=sh.id, sequence_order=i))

    # 人员组别
    eg_stamp = EmployeeGroup(name="冲压组", description="冲压岗")
    eg_asm = EmployeeGroup(name="装配组", description="装配岗")
    db.add_all([eg_stamp, eg_asm]); db.flush()

    # 员工 + 技能矩阵
    def _emp(code, name, eg, shift, devices, procs, rotation=None, start_shift=None, start_date=None):
        e = Employee(name=name, group_id=eg.id, employee_code=code, shift_id=shift.id,
                     rotation_id=rotation.id if rotation else None,
                     rotation_start_shift_id=start_shift.id if start_shift else None,
                     rotation_start_date=start_date)
        db.add(e); db.flush()
        for d in devices:
            db.execute(employee_device.insert().values(employee_id=e.id, device_id=d.id))
        for pr in procs:
            db.execute(employee_process.insert().values(employee_id=e.id, process_id=pr.id))
        return e
    _emp("E001", "张三", eg_stamp, sh_morning, [d1, d2], [pro_stamp, pro_drill])
    _emp("E002", "李四", eg_asm, sh_morning, [d3, d4], [pro_asm])
    _emp("E003", "王五", eg_stamp, sh_morning, [d1, d2], [pro_stamp, pro_drill])
    _emp("E004", "赵六", eg_asm, sh_morning, [d3, d4], [pro_asm])
    _emp("E005", "钱七", eg_stamp, sh_morning, [d1], [pro_stamp], rotation=rot, start_shift=sh_morning, start_date=datetime.now().strftime('%Y-%m-%d'))

    db.commit()
    return {"ok": True, "message": "示例数据载入完成：3个产品、4台设备、3道工序、3条工艺路线、5名员工"}


if __name__ == "__main__":
    import uvicorn, argparse, os, secrets
    from backend.models.database import SessionLocal

    parser = argparse.ArgumentParser(description='生产排程系统服务器')
    parser.add_argument('--port', type=int, default=8899, help='监听端口 (默认 8899)')
    parser.add_argument('--host', type=str, default='0.0.0.0', help='监听地址 (默认 0.0.0.0)')
    parser.add_argument('--admin-password', type=str, default=None,
                        help='首次启动创建的默认管理员密码，默认自动随机生成')
    args = parser.parse_args()

    # 首次启动时创建默认管理员（仅在系统无任何用户时）
    db = SessionLocal()
    if not db.query(User).first():
        default_pwd = os.environ.get('ADMIN_PASSWORD') or args.admin_password or secrets.token_urlsafe(12)
        admin = User(username='admin', password_hash=get_password_hash(default_pwd), role='admin')
        db.add(admin)
        db.commit()
        print(f"[提示] 已创建默认管理员账号 admin")
        print(f"[提示] 初始密码: {default_pwd}  (请登录后尽快修改，或通过环境变量 ADMIN_PASSWORD 预设)")
    db.close()

    print(f"启动生产排程系统服务器 http://{args.host}:{args.port}")
    uvicorn.run(app, host=args.host, port=args.port)
